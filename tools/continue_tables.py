"""Split tables only where Word actually broke them across pages; insert «Продолжение Таблицы N»."""
from __future__ import annotations

import argparse
import json
import os
import re
import subprocess
import sys
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from tools.paths import assert_not_sample
from tools.render_docx import word_available
from tools.replace_body import _restore_bottom_border


class ContinueError(Exception):
    pass


_CAP_N = re.compile(r"(?:Таблица|Таблицы)\s+(\d+)")
WD_ACTIVE_END_PAGE = 3
WD_COLLAPSE_START = 1


def _caption_n(text: str) -> int | None:
    m = _CAP_N.search(text or "")
    return int(m.group(1)) if m else None


def _row_page(row) -> int:
    """Страница начала строки. Конец строки часто попадает на следующую и даёт ложный разрыв."""
    try:
        rng = row.Cells(1).Range
    except Exception:
        rng = row.Range
    rng.Collapse(WD_COLLAPSE_START)
    return int(rng.Information(WD_ACTIVE_END_PAGE))


def _collect_splits_word_impl(docx_path: Path) -> list[dict]:
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    doc = None
    splits: list[dict] = []
    try:
        doc = word.Documents.Open(str(docx_path.resolve()), False, True, False)
        try:
            doc.Repaginate()
        except Exception:
            pass
        for ti in range(1, doc.Tables.Count + 1):
            tbl = doc.Tables(ti)
            if tbl.Columns.Count <= 1 or tbl.Rows.Count < 3:
                continue
            pages = [_row_page(tbl.Rows(ri)) for ri in range(1, tbl.Rows.Count + 1)]
            first = pages[0]
            split_at = next((i + 1 for i, p in enumerate(pages) if p != first), None)
            # Need header + at least one data row on the first page, otherwise
            # the first chunk would be a header-only stub.
            if split_at is None or split_at <= 2:
                continue
            pos = int(tbl.Range.Start)
            chunk = doc.Range(max(0, pos - 500), pos).Text
            n = _caption_n(chunk) or ti
            splits.append({"ti": ti - 1, "split_row": split_at - 1, "n": n})
    finally:
        if doc is not None:
            try:
                doc.Close(False)
            except Exception:
                pass
        try:
            word.Quit()
        except Exception:
            pass
        pythoncom.CoUninitialize()
    return splits


def _continuation_paragraph(n: int):
    p = OxmlElement("w:p")
    pPr = OxmlElement("w:pPr")
    pPr.append(OxmlElement("w:keepNext"))
    jc = OxmlElement("w:jc")
    jc.set(qn("w:val"), "left")
    pPr.append(jc)
    p.append(pPr)
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rPr.append(OxmlElement("w:i"))
    rPr.append(OxmlElement("w:iCs"))
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "24")
    rPr.append(sz)
    r.append(rPr)
    t = OxmlElement("w:t")
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = f"Продолжение Таблицы {n}"
    r.append(t)
    p.append(r)
    return p


def _split_tbl_at(tbl, split_row: int, n: int) -> None:
    trs = tbl.findall(qn("w:tr"))
    if split_row <= 1 or split_row >= len(trs):
        return
    new_tbl = deepcopy(tbl)
    new_trs = new_tbl.findall(qn("w:tr"))
    for tr in list(new_trs[1:split_row]):
        new_tbl.remove(tr)
    for tr in list(trs[split_row:]):
        tbl.remove(tr)
    # Оба куска — закрытые таблицы: нижняя линия есть и у части, которая осталась на странице.
    _restore_bottom_border(tbl)
    _restore_bottom_border(new_tbl)
    cap = _continuation_paragraph(n)
    tbl.addnext(new_tbl)
    tbl.addnext(cap)


def _merge_pair(prev_tbl, next_tbl) -> bool:
    cap = prev_tbl.getnext()
    if cap is None or cap.tag != qn("w:p"):
        return False
    text = "".join(t.text or "" for t in cap.findall(".//" + qn("w:t"))).strip()
    if not text.startswith("Продолжение"):
        return False
    if cap.getnext() is not next_tbl:
        return False
    trs = next_tbl.findall(qn("w:tr"))
    # Повторенная шапка продолжения не переносится второй раз.
    _restore_bottom_border(prev_tbl)
    for tr in list(trs[1:]):
        next_tbl.remove(tr)
        prev_tbl.append(tr)
    parent = next_tbl.getparent()
    if parent is None:
        return False
    parent.remove(cap)
    parent.remove(next_tbl)
    _restore_bottom_border(prev_tbl)
    return True


def _collect_merges_word_impl(docx_path: Path) -> list[dict]:
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    doc = None
    merges: list[dict] = []
    try:
        doc = word.Documents.Open(str(docx_path.resolve()), False, True, False)
        try:
            doc.Repaginate()
        except Exception:
            pass
        count = int(doc.Tables.Count)
        for ti in range(1, count):
            prev = doc.Tables(ti)
            nxt = doc.Tables(ti + 1)
            gap = doc.Range(int(prev.Range.End), int(nxt.Range.Start)).Text or ""
            if "Продолжение" not in gap:
                continue
            plain = gap.replace("\r", "").replace("\x07", "").replace("\x0b", "").strip()
            if len(plain) > 80:
                continue
            try:
                prev_page = _row_page(prev.Rows(prev.Rows.Count))
                next_page = _row_page(nxt.Rows(1))
            except Exception:
                continue
            if prev_page == next_page:
                merges.append({"prev": ti - 1, "next": ti})
    finally:
        if doc is not None:
            try:
                doc.Close(False)
            except Exception:
                pass
        try:
            word.Quit()
        except Exception:
            pass
        pythoncom.CoUninitialize()
    return merges


def _apply_merges(docx_path: Path, merges: list[dict]) -> int:
    if not merges:
        return 0
    document = Document(str(docx_path))
    tables = list(document.tables)
    applied = 0
    for spec in sorted(merges, key=lambda s: s["next"], reverse=True):
        prev_i, next_i = spec["prev"], spec["next"]
        if prev_i < 0 or next_i >= len(tables) or next_i != prev_i + 1:
            continue
        if _merge_pair(tables[prev_i]._tbl, tables[next_i]._tbl):
            applied += 1
    if applied:
        document.save(str(docx_path))
    return applied


def _collect_merges(docx_path: Path) -> list[dict]:
    env = os.environ.copy()
    env["_REPORT_WORD_WORKER"] = "1"
    proc = subprocess.run(
        [sys.executable, "-m", "tools.continue_tables", "--docx", str(docx_path), "--collect-merges"],
        env=env,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    if proc.returncode != 0:
        raise ContinueError((proc.stderr or proc.stdout or "Word не смог сверить продолжения таблиц").strip())
    raw = (proc.stdout or "").strip()
    try:
        return json.loads(raw) if raw else []
    except json.JSONDecodeError as exc:
        raise ContinueError(f"Не разобрать склейку таблиц: {exc}") from exc


def _apply_splits(docx_path: Path, splits: list[dict]) -> int:
    if not splits:
        return 0
    document = Document(str(docx_path))
    tables = list(document.tables)
    applied = 0
    for spec in sorted(splits, key=lambda s: s["ti"], reverse=True):
        ti = spec["ti"]
        if ti < 0 or ti >= len(tables):
            continue
        _split_tbl_at(tables[ti]._tbl, spec["split_row"], spec["n"])
        applied += 1
    document.save(str(docx_path))
    return applied


def _collect_splits(docx_path: Path) -> list[dict]:
    env = os.environ.copy()
    env["_REPORT_WORD_WORKER"] = "1"
    proc = subprocess.run(
        [sys.executable, "-m", "tools.continue_tables", "--docx", str(docx_path), "--collect"],
        env=env,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    if proc.returncode != 0:
        raise ContinueError((proc.stderr or proc.stdout or "Word не смог разметить продолжения таблиц").strip())
    raw = (proc.stdout or "").strip()
    try:
        return json.loads(raw) if raw else []
    except json.JSONDecodeError as exc:
        raise ContinueError(f"Не разобрать разметку таблиц: {exc}") from exc


def continue_tables(docx_path: Path) -> int:
    docx_path = Path(docx_path)
    assert_not_sample(docx_path)
    if not word_available():
        return 0
    total = 0
    for _ in range(6):
        splits = _collect_splits(docx_path)
        if not splits:
            break
        applied = _apply_splits(docx_path, splits)
        if applied <= 0:
            break
        total += applied
    # Ложный разрыв оставляет два «Продолжение» на одной странице: такие куски склеиваются.
    for _ in range(6):
        merges = _collect_merges(docx_path)
        if not merges:
            break
        if _apply_merges(docx_path, merges) <= 0:
            break
    return total


def main() -> int:
    p = argparse.ArgumentParser()
    p.add_argument("--docx", required=True)
    p.add_argument("--collect", action="store_true")
    p.add_argument("--collect-merges", action="store_true")
    args = p.parse_args()
    try:
        if args.collect:
            print(json.dumps(_collect_splits_word_impl(Path(args.docx)), ensure_ascii=False))
            return 0
        if args.collect_merges:
            print(json.dumps(_collect_merges_word_impl(Path(args.docx)), ensure_ascii=False))
            return 0
        n = continue_tables(Path(args.docx))
        print(f"продолжений: {n}")
    except ContinueError as exc:
        print(str(exc), file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
