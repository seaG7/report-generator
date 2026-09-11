from __future__ import annotations

import argparse
import os
import shutil
import sys
from pathlib import Path

from docx import Document

from tools.config import Student, TitleFields, WorkMeta, load_student, load_title_fields, load_work, resolve_to
from tools.paths import assert_not_sample


class TitleFieldNotFound(Exception):
    pass


class TitleNotConfirmed(Exception):
    pass


def _iter_paragraphs(document: Document):
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in document.sections:
        yield from section.header.paragraphs
        yield from section.footer.paragraphs
        if section.different_first_page_header_footer:
            yield from section.first_page_header.paragraphs
            yield from section.first_page_footer.paragraphs


def _paragraph_has(paragraph, needle: str) -> bool:
    return needle in (paragraph.text or "")


def _replace_in_paragraph(paragraph, old: str, new: str) -> bool:
    if old not in (paragraph.text or ""):
        return False
    text = paragraph.text.replace(old, new)
    if not paragraph.runs:
        paragraph.add_run(text)
        return True
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""
    return True


def _collect_all_text(document: Document) -> str:
    return "\n".join(p.text or "" for p in _iter_paragraphs(document))


def _apply_mapping(docx_path: Path, mapping: dict[str, str]) -> None:
    document = Document(str(docx_path))
    blob = _collect_all_text(document)
    for find in mapping:
        if find not in blob:
            raise TitleFieldNotFound(f"Строка не найдена на титульнике: {find}")
    for paragraph in _iter_paragraphs(document):
        for old, new in mapping.items():
            _replace_in_paragraph(paragraph, old, new)
    document.save(str(docx_path))


def patch_title(
    docx_path: Path,
    title_fields: TitleFields,
    student: Student,
    work: WorkMeta,
) -> None:
    docx_path = Path(docx_path)
    assert_not_sample(docx_path)
    if not title_fields.confirmed:
        raise TitleNotConfirmed("title-fields.yaml не подтверждён")
    mapping = {
        field.find: resolve_to(field.to, student, work)
        for field in title_fields.fields.values()
    }
    tmp = docx_path.with_suffix(".tmp.docx")
    shutil.copy2(docx_path, tmp)
    try:
        _apply_mapping(tmp, mapping)
        os.replace(tmp, docx_path)
    except Exception:
        if tmp.exists():
            tmp.unlink()
        raise


def main() -> int:
    p = argparse.ArgumentParser()
    p.add_argument("--docx", required=True)
    p.add_argument("--title-fields", required=True)
    p.add_argument("--student", required=True)
    p.add_argument("--work", required=True)
    args = p.parse_args()
    try:
        patch_title(
            Path(args.docx),
            load_title_fields(Path(args.title_fields)),
            load_student(Path(args.student)),
            load_work(Path(args.work)),
        )
    except (TitleFieldNotFound, TitleNotConfirmed) as exc:
        print(str(exc), file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
