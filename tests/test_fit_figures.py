from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from tests.test_replace_body import TF, _make_work
from tools.fit_figures import _apply_fits
from tools.insert_figure import scale_drawing_to_height
from tools.replace_body import replace_body


def test_scale_drawing_keeps_aspect(tmp_path: Path):
    work, docx, md = _make_work(
        tmp_path,
        "# Цель работы\n\nАбзац перед рисунком.\n\n![Первый](screenshots/01-one.png)\n",
    )
    replace_body(docx, md, TF, work)
    doc = Document(str(docx))
    fig = None
    for p in doc.paragraphs:
        if p._element.find(".//" + qn("w:drawing")) is not None:
            cap = p._element.getnext()
            if cap is not None and (cap.xpath("string()") or "").startswith("Рисунок"):
                fig = p._element
                break
    assert fig is not None
    ext = fig.find(".//" + "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent")
    cx0, cy0 = int(ext.get("cx")), int(ext.get("cy"))
    target = max(1, cy0 // 2)
    assert scale_drawing_to_height(fig, target)
    cx1, cy1 = int(ext.get("cx")), int(ext.get("cy"))
    assert cy1 == target
    assert abs(cx1 / cx0 - cy1 / cy0) < 0.02


def test_apply_fits_shrinks_named_figure(tmp_path: Path):
    work, docx, md = _make_work(
        tmp_path,
        "# Цель работы\n\nАбзац перед рисунком.\n\n![Первый](screenshots/01-one.png)\n",
    )
    replace_body(docx, md, TF, work)
    doc = Document(str(docx))
    ext = None
    for p in doc.paragraphs:
        drawing = p._element.find(".//" + qn("w:drawing"))
        if drawing is None:
            continue
        cap = p._element.getnext()
        if cap is not None and (cap.xpath("string()") or "").startswith("Рисунок 1"):
            ext = p._element.find(
                ".//" + "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent"
            )
            break
    assert ext is not None
    cy0 = int(ext.get("cy"))
    target = max(1, cy0 // 2)
    assert _apply_fits(docx, [{"n": 1, "cy": target}]) == 1
    out = Document(str(docx))
    for p in out.paragraphs:
        if p._element.find(".//" + qn("w:drawing")) is None:
            continue
        cap = p._element.getnext()
        if cap is None or not (cap.xpath("string()") or "").startswith("Рисунок 1"):
            continue
        ext = p._element.find(
            ".//" + "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent"
        )
        assert int(ext.get("cy")) == target
        return
    raise AssertionError("рисунок 1 не найден")
