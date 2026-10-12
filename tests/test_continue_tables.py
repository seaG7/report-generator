from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from lxml import etree

from tools.continue_tables import _apply_splits, _merge_pair


def test_split_inserts_caption_as_paragraph_not_cell(tmp_path: Path):
    doc = Document()
    doc.add_paragraph("Таблица 5 — Теги")
    tbl = doc.add_table(rows=5, cols=3)
    for i, row in enumerate(tbl.rows):
        row.cells[0].text = "Страница" if i == 0 else f"row-{i}"
        row.cells[1].text = "Title" if i == 0 else "t"
        row.cells[2].text = "H2" if i == 0 else "h"
    path = tmp_path / "t.docx"
    doc.save(path)

    applied = _apply_splits(path, [{"ti": 0, "split_row": 3, "n": 5}])
    assert applied == 1

    out = Document(str(path))
    assert len(out.tables) == 2
    assert [c.text for c in out.tables[0].rows[0].cells] == ["Страница", "Title", "H2"]
    assert len(out.tables[0].rows) == 3
    assert len(out.tables[1].rows) == 3
    assert out.tables[1].rows[0].cells[0].text == "Страница"
    assert out.tables[1].rows[1].cells[0].text == "row-3"

    for tbl in out.tables:
        for row in tbl.rows:
            for cell in row.cells:
                assert "Продолжение" not in (cell.text or "")

    texts = [p.text.strip() for p in out.paragraphs]
    assert "Продолжение Таблицы 5" in texts
    cap = next(p for p in out.paragraphs if p.text.strip() == "Продолжение Таблицы 5")
    rPr = cap.runs[0]._element.find(qn("w:rPr"))
    assert rPr is not None and rPr.find(qn("w:i")) is not None
    jc = cap._element.find(qn("w:pPr")).find(qn("w:jc"))
    assert jc is not None and jc.get(qn("w:val")) == "left"

    tags = [etree.QName(child).localname for child in out.element.body]
    assert tags.count("tbl") == 2
    first_tbl = tags.index("tbl")
    assert tags[first_tbl + 1] == "p"
    assert tags[first_tbl + 2] == "tbl"
    assert out.element.body[first_tbl + 1].xpath("string()") == "Продолжение Таблицы 5"
    bottom = out.tables[0]._tbl.find(qn("w:tblPr")).find(qn("w:tblBorders")).find(qn("w:bottom"))
    assert bottom is not None and bottom.get(qn("w:val")) == "single"
    last_cell = out.tables[0].rows[-1].cells[0]._tc
    cell_bottom = last_cell.find(qn("w:tcPr")).find(qn("w:tcBorders")).find(qn("w:bottom"))
    assert cell_bottom is not None and cell_bottom.get(qn("w:val")) == "single"


def test_same_page_continuations_merge_back(tmp_path: Path):
    doc = Document()
    doc.add_paragraph("Таблица 11 — Частотность")
    first = doc.add_table(rows=3, cols=2)
    first.rows[0].cells[0].text = "Запрос"
    first.rows[1].cells[0].text = "подобно"
    first.rows[2].cells[0].text = "выразить"
    cap = doc.add_paragraph("Продолжение Таблицы 11")
    second = doc.add_table(rows=3, cols=2)
    second.rows[0].cells[0].text = "Запрос"
    second.rows[1].cells[0].text = "подстрока"
    second.rows[2].cells[0].text = "естьnull"
    assert _merge_pair(first._tbl, second._tbl)
    assert cap._element.getparent() is None
    assert len(doc.tables) == 1
    texts = [row.cells[0].text for row in doc.tables[0].rows]
    assert texts == ["Запрос", "подобно", "выразить", "подстрока", "естьnull"]
    assert not any(p.text.startswith("Продолжение") for p in doc.paragraphs)
