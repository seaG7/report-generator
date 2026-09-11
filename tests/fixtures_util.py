from pathlib import Path
from PIL import Image
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas


def make_two_image_pdf(path: Path) -> None:
    path = Path(path)
    img_dir = path.parent / "_pdf_imgs"
    img_dir.mkdir(parents=True, exist_ok=True)
    red = img_dir / "red.png"
    blue = img_dir / "blue.png"
    Image.new("RGB", (40, 40), (255, 0, 0)).save(red)
    Image.new("RGB", (40, 40), (0, 0, 255)).save(blue)
    c = canvas.Canvas(str(path), pagesize=A4)
    w, h = A4
    c.setFont("Times-Roman", 14)
    c.drawString(72, h - 72, "BEFORE")
    c.drawImage(str(red), 72, h - 160, 40, 40)
    c.drawString(72, h - 200, "MIDDLE")
    c.drawImage(str(blue), 72, h - 280, 40, 40)
    c.drawString(72, h - 320, "AFTER")
    c.save()


def make_sample_docx(path: Path, image: Path | None = None) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    if image is None:
        image = path.parent / "_sample_fig.png"
        Image.new("RGB", (40, 40), (180, 180, 180)).save(image)
    doc = Document()
    doc.add_paragraph("Университет")
    doc.add_paragraph("Петров Пётр Петрович")
    doc.add_paragraph("ИКБО-00-00")
    doc.add_paragraph("Лабораторная работа №1. Пример")
    doc.add_page_break()
    doc.add_heading("Цель работы", level=1)
    doc.add_paragraph("Текст цели.")
    p = doc.add_paragraph()
    p.add_run().add_picture(str(image), width=Inches(1.5))
    cap = doc.add_paragraph("Рисунок 1 – Пример")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("Таблица 1 – Пример")
    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    table.cell(1, 0).text = "1"
    table.cell(1, 1).text = "2"
    doc.add_heading("Выводы", level=1)
    doc.add_paragraph("Текст выводов.")
    doc.save(str(path))
